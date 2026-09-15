"""Exercise archive, country and request-shape failure modes with workbook fixtures."""
import importlib.util
import io
from datetime import datetime
from pathlib import Path
import unittest
from unittest.mock import patch
import openpyxl

spec = importlib.util.spec_from_file_location('rigs', Path(__file__).resolve().parents[1] / 'instrument_check.py')
rigs = importlib.util.module_from_spec(spec)
spec.loader.exec_module(rigs)
UUID = 'ac9a8c23-3a35-4a0c-88d3-d8678af2f5c6'
LABEL = 'North America Rig Count Report - New Report'
QUERY = 'https://rigcount.bakerhughes.com/na-rig-count|' + LABEL


class Clock(datetime):
    @classmethod
    def now(cls):
        return cls(2026, 9, 15)


class Response(io.BytesIO):
    status = 200

    def __init__(self, blob, disposition=''):
        super().__init__(blob)
        self.headers = {'Content-Disposition': disposition}


def workbook(us_oil=450, stamp='11/09/2026', duplicate=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'NAM Summary'
    ws['D4'] = stamp
    ws['B19'] = 'U.S. Breakout Information'
    ws['D19'], ws['E19'], ws['F19'] = 'This Week', '+/-', 'Last Week'
    if us_oil is not None:
        ws['B22'], ws['D22'] = 'Oil', us_oil
    if duplicate:
        ws['B23'], ws['D23'] = 'Oil', 451
    ws['B29'] = 'Canada Breakout Information'
    ws['B32'], ws['D32'] = 'Oil', 141
    blob = io.BytesIO()
    wb.save(blob)
    return blob.getvalue()


def changed_cells(**values):
    wb = openpyxl.load_workbook(io.BytesIO(workbook()))
    ws = wb['NAM Summary']
    for address, value in values.items():
        ws[address] = value
    blob = io.BytesIO()
    wb.save(blob)
    return blob.getvalue()


class RigReader(unittest.TestCase):
    def run_probe(self, blob=None, filename='09-11-2026 report.xlsx', extra_anchor='', fail_first=False):
        html = f'<a href="/static-files/{UUID}">{LABEL}</a>{extra_anchor}'.encode()
        results = [Response(html), Response(blob if blob is not None else workbook(), filename)]
        if fail_first:
            results.insert(0, OSError('request timeout'))
        with patch.object(rigs, 'datetime', Clock):
            with patch.object(rigs.urllib.request, 'urlopen', side_effect=results) as get:
                result = rigs.probe_bhrigs(QUERY)
                return result, get.call_args_list

    def test_primary_us_oil_and_minimal_request(self):
        result, calls = self.run_probe()
        self.assertTrue(result[0])
        self.assertIn('US OIL rig count 450', result[2])
        self.assertEqual(len(calls), 2)
        self.assertEqual(calls[0].args[0].get_header('User-agent'), 'Mozilla/5.0')

    def test_bounded_fallback_after_transport_failure(self):
        result, calls = self.run_probe(fail_first=True)
        self.assertTrue(result[0])
        self.assertEqual(len(calls), 3)
        self.assertEqual(calls[1].args[0].get_header('User-agent'), rigs.BROWSER_UA)

    def test_archive_decoy_fails(self):
        result, _ = self.run_probe(filename='08-29-2025 report.xlsx')
        self.assertFalse(result[0])
        self.assertIn('2025-08-29', result[2])

    def test_filename_cannot_refresh_old_workbook(self):
        result, _ = self.run_probe(workbook(stamp='04/09/2026'))
        self.assertFalse(result[0])
        self.assertIn('disagrees', result[2])

    def test_missing_us_oil_does_not_use_canada(self):
        result, _ = self.run_probe(workbook(us_oil=None))
        self.assertFalse(result[0])

    def test_fractional_and_duplicate_oil_fail(self):
        for blob in [workbook(us_oil=450.5), workbook(duplicate=True)]:
            self.assertFalse(self.run_probe(blob)[0][0])

    def test_same_label_different_download_is_ambiguous(self):
        extra = f'<a href="/static-files/00000000-0000-0000-0000-000000000000">{LABEL}</a>'
        result, calls = self.run_probe(extra_anchor=extra)
        self.assertFalse(result[0])
        self.assertEqual(len(calls), 1)

    def test_blank_current_never_uses_change_or_prior_count(self):
        for values in [dict(D22=None, E22=1, F22=449),
                       dict(D22=None, F22=450), dict(D22=None)]:
            with self.subTest(values=values):
                result, _ = self.run_probe(changed_cells(**values))
                self.assertFalse(result[0])
                self.assertIn('invalid US Oil count None', result[2])

    def test_missing_or_duplicate_current_header_fails(self):
        for values in [dict(D19=None), dict(F19='This Week')]:
            with self.subTest(values=values):
                result, _ = self.run_probe(changed_cells(**values))
                self.assertFalse(result[0])
                self.assertIn('exactly one This Week', result[2])

    def test_current_column_follows_header_not_old_offset(self):
        blob = changed_cells(D19='Last Week', F19='This Week',
                             D22=450, E22=10, F22=460)
        result, _ = self.run_probe(blob)
        self.assertTrue(result[0])
        self.assertIn('US OIL rig count 460', result[2])

    def test_invalid_current_values_fail_even_with_valid_prior(self):
        for value in [True, '450', -1, 450.5]:
            with self.subTest(value=value):
                result, _ = self.run_probe(changed_cells(D22=value, F22=449))
                self.assertFalse(result[0])
                self.assertIn('invalid US Oil count', result[2])

    def test_duplicate_us_header_fails(self):
        result, _ = self.run_probe(changed_cells(
            B20='U.S. Breakout Information', D20='This Week'))
        self.assertFalse(result[0])
        self.assertIn('duplicate U.S. Breakout', result[2])

    def test_real_zero_current_count_is_not_missing(self):
        result, _ = self.run_probe(changed_cells(D22=0, F22=449))
        self.assertTrue(result[0])
        self.assertIn('US OIL rig count 0 ', result[2])


if __name__ == '__main__':
    unittest.main()
