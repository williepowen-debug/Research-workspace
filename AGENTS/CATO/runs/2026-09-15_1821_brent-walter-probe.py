"""Read-only review reproduction; mocked network, saved publisher row layout.

Run from repo root with .venv/bin/python3. No production repair or live grade.
"""
import sys
sys.dont_write_bytecode = True
from datetime import datetime
import importlib.util
import io
import json
from pathlib import Path
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
spec = importlib.util.spec_from_file_location(
    'reviewed_rigs', ROOT / 'AGENTS/BRENT/scripts/instrument_check.py')
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)
source = openpyxl.load_workbook(
    ROOT / 'AGENTS/BRENT/research/2026-09-15_workdown/bh-current.xlsx',
    read_only=True, data_only=True)


class Clock(datetime):
    @classmethod
    def now(cls):
        return cls(2026, 9, 15)


class Response(io.BytesIO):
    status = 200

    def __init__(self, raw, disposition=''):
        super().__init__(raw)
        self.headers = {'Content-Disposition': disposition}


def run_case(blank_current):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'NAM Summary'
    for row in source['NAM Summary'].iter_rows(max_row=60, values_only=True):
        sheet.append(row)
    if blank_current:
        sheet['D22'] = None
    blob = io.BytesIO()
    book.save(blob)
    label = 'North America Rig Count Report - New Report'
    html = ('<a href="/static-files/ac9a8c23-3a35-4a0c-88d3-d8678af2f5c6">'
            + label + '</a>').encode()
    replies = [Response(html), Response(blob.getvalue(), '09-11-2026 report.xlsx')]
    with patch.object(reader, 'datetime', Clock), patch.object(
            reader.urllib.request, 'urlopen', side_effect=replies):
        ok, stamp, detail = reader.probe_bhrigs(
            'https://rigcount.bakerhughes.com/na-rig-count|' + label)
    return {'blank_current': blank_current, 'success': ok,
            'date': stamp.isoformat() if stamp else None, 'detail': detail}


if __name__ == '__main__':
    results = [run_case(False), run_case(True)]
    source.close()
    print(json.dumps(results, indent=2))
    assert results[0]['success'] and 'US OIL rig count 450 ' in results[0]['detail']
    assert results[1]['success'] and 'US OIL rig count 1 ' in results[1]['detail']
    print('REPRODUCED: blank This Week cell is replaced by the +1 change column.')
