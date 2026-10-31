#!/usr/bin/env python3
"""L616 breadth — build the constituent universe from issuer holdings files.

Inputs (data only, never imported): downloads/holdings-daily-us-en-{spy,mdy,spsm,xl*}.xlsx
  SSGA daily holdings, all stamped "As of 29-Oct-2026" (the book entering 10/30).
Output: universe.csv — ticker, yf_ticker, name, index, fund_weight_pct, shares_held, gics_sector
  S&P 500 sector = membership in the 11 Select Sector SPDR files (GICS by construction);
  S&P 400 sector = MDY file's own Sector column; S&P 600 sector = not in SPSM file (blank).
Run: .venv/bin/python3 -I <this> <research_dir>
"""
import sys, os, csv
import openpyxl

R = sys.argv[1]
D = os.path.join(R, "downloads")
SECT = {"xlb": "Materials", "xlc": "Communication Services", "xle": "Energy",
        "xlf": "Financials", "xli": "Industrials", "xlk": "Information Technology",
        "xlp": "Consumer Staples", "xlre": "Real Estate", "xlu": "Utilities",
        "xlv": "Health Care", "xly": "Consumer Discretionary"}
NONEQ = {"CASH_USD", "-", "", None}


def read(fund):
    wb = openpyxl.load_workbook(os.path.join(D, f"holdings-daily-us-en-{fund}.xlsx"),
                                read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    asof = rows[2][1]
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "Name" and r[1] == "Ticker")
    hdr = rows[hdr_i]
    out = []
    for r in rows[hdr_i + 1:]:
        if not r or r[1] in NONEQ or r[0] is None:
            continue
        d = dict(zip(hdr, r))
        if str(d.get("Sector")) == "Unassigned":
            continue
        # non-equity rows: acquisition contra entries, earnout rights, index futures legs,
        # numeric placeholder identifiers (2602335D-style) — not constituents, excluded by rule
        nm = str(d.get("Name") or "").upper()
        tk = str(d.get("Ticker") or "").strip()
        if any(k in nm for k in ("CONTRA ", "EARNOUT", "E-MINI", "FUTURE")) or (tk[:-1].isdigit() and tk.endswith("D")):
            continue
        out.append(d)
    return asof, out


def yf(t):
    return str(t).strip().replace(".", "-").replace("/", "-")


sector_of = {}
asofs = {}
for f, s in SECT.items():
    a, rows = read(f)
    asofs[f] = a
    for d in rows:
        sector_of.setdefault(str(d["Ticker"]).strip(), s)

uni = []
for fund, idx in (("spy", "SP500"), ("mdy", "SP400"), ("spsm", "SP600")):
    a, rows = read(fund)
    asofs[fund] = a
    for d in rows:
        t = str(d["Ticker"]).strip()
        sec = sector_of.get(t, "") if idx == "SP500" else ("" if str(d.get("Sector")) in ("-", "None") else str(d.get("Sector")))
        uni.append({"ticker": t, "yf_ticker": yf(t), "name": d["Name"], "index": idx,
                    "fund_weight_pct": d["Weight"], "shares_held": d["Shares Held"],
                    "gics_sector": sec})

# a name held in two funds keeps its first (largest-index) membership only
seen, dedup = set(), []
for u in uni:
    if u["yf_ticker"] in seen:
        continue
    seen.add(u["yf_ticker"])
    dedup.append(u)

with open(os.path.join(R, "universe.csv"), "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(dedup[0].keys()))
    w.writeheader()
    w.writerows(dedup)

from collections import Counter
print("as-of stamps:", sorted(set(asofs.values())))
print("rows:", Counter(u["index"] for u in dedup), "dropped dups:", len(uni) - len(dedup))
sp5 = [u for u in dedup if u["index"] == "SP500"]
print("SP500 with sector:", sum(1 for u in sp5 if u["gics_sector"]), "of", len(sp5))
print("SP500 unmapped:", [u["ticker"] for u in sp5 if not u["gics_sector"]][:20])
print("SP500 sector counts:", Counter(u["gics_sector"] for u in sp5))
print("SP500 fund weight sum:", round(sum(float(u["fund_weight_pct"] or 0) for u in sp5), 3))
